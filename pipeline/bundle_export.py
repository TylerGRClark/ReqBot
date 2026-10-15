"""One workbook for a bundle: a Contents sheet first, then one checklist sheet per publication that is loaded."""
import io
import re
from pathlib import Path

from core.profiles import load_profile
from pipeline import checklist_export
from services import checklist_service

_BAD_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")


def _tab_title(key: str, used: set) -> str:
    base = (_BAD_TITLE_CHARS.sub("_", key).strip("' ") or "sheet")[:31]
    title, n = base, 2
    while title.lower() in used or title.lower() == "contents":
        suffix = f"_{n}"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def build_bundle_workbook(processed_dir: Path, bundle: dict, profile: str = "cybersecurity") -> tuple[bytes, list[dict]]:
    """Returns (xlsx bytes, one summary dict per listed publication: doc, title, tab, rows, with_question, loaded).

    An unknown profile raises (FileNotFoundError or ValueError) instead of showing every publication as "not loaded"."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    load_profile(profile)
    wb = Workbook()
    contents = wb.active
    contents.title = "Contents"
    used: set = set()
    summary = []
    for entry in bundle["documents"]:
        key, title = entry["doc"], entry.get("title") or ""
        rec = {"doc": key, "title": title, "tab": "", "rows": 0, "with_question": 0, "loaded": False}
        try:
            checklist = checklist_service.generate(Path(processed_dir), key, profile)
        except (ValueError, FileNotFoundError):
            checklist = None
        if checklist is not None:
            items = checklist.get("items", [])
            rec.update(
                tab=_tab_title(key, used), rows=len(items), loaded=True,
                with_question=sum(1 for i in items if (i.get("audit_question") or "").strip()),
            )
            checklist_export.to_xlsx(checklist, workbook=wb, sheet_title=rec["tab"])
        summary.append(rec)

    contents.sheet_view.showGridLines = False
    contents.cell(row=1, column=1, value=checklist_export._csv_safe(bundle.get("label") or "Checklist bundle")).font = Font(name="Calibri", size=16, bold=True)
    note = bundle.get("note") or "A starting list. Add or remove publications to match the unit's real programs."
    contents.cell(row=2, column=1, value=checklist_export._csv_safe(note)).font = Font(name="Calibri", size=10, color="4B5563")
    headers = ["Tab", "Publication", "Rows", "With a draft question", "In this workbook"]
    for col, h in enumerate(headers, start=1):
        cell = contents.cell(row=4, column=col, value=h)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="047857")
    for n, rec in enumerate(summary, start=5):
        values = [rec["tab"] or "", rec["title"] or rec["doc"], rec["rows"] if rec["loaded"] else "",
                  rec["with_question"] if rec["loaded"] else "", "yes" if rec["loaded"] else "not loaded yet"]
        for col, v in enumerate(values, start=1):
            contents.cell(row=n, column=col, value=checklist_export._csv_safe(v)).alignment = Alignment(vertical="top", wrap_text=col == 2)
    for letter, width in zip("ABCDE", (18, 70, 8, 22, 18)):
        contents.column_dimensions[letter].width = width
    contents.freeze_panes = "A5"
    wb.properties.title = str(bundle.get("label") or "Checklist bundle")[:200]
    wb.properties.creator = "ReqBot"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary

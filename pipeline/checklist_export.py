"""Checklist export — serialize a checklist envelope dict to CSV, JSON, Markdown, or XLSX.

Input:  the dict returned by services/checklist_service.py::generate()
Output: a string (CSV/JSON/Markdown) or bytes (XLSX); caller writes to file or stdout.

All display and export logic lives here and in cli/reqbot.py (WP-21.5).
"""
import csv
import io
import json
from datetime import datetime

# CSV column order: locate → ask → record → verify → trace
_CSV_COLUMNS = [
    "source_ref",
    "citation",
    "section_title_path",
    "section_heading",
    "page_refs",
    "applies_to",
    "parent_paragraph",
    "source_quote",
    "passage",
    "audit_question",
    "status",
    "assessor_notes",
    "item_flags",
    "requires_human_review",
    "review_reasons",
    "confidence",
    "checklist_item_id",
    "requirement_ids",
    "domain_tags",
]


_FORMULA_CHARS = frozenset("=+-@")


def _join(values: list, sep: str) -> str:
    return sep.join(str(v) for v in values)


def _csv_safe(value: object) -> object:
    """Prefix formula-like string cells with a single quote to block spreadsheet injection.

    Excel and LibreOffice treat cells whose effective first character is =, +, -, or @
    as formulas regardless of CSV quoting. Prefixing with ' is the standard mitigation.
    Non-string values (bool, float) are returned unchanged.
    """
    if isinstance(value, str):
        stripped = value.lstrip()
        if stripped and stripped[0] in _FORMULA_CHARS:
            return "'" + value
    return value


def _needs_attention(item: dict) -> bool:
    """A row worth a reader's second look in the sheet: a specific hint (item_flags) or a missing citation/tag. Confidence alone no longer shades every row (the pipeline's
    confidence is below the review threshold on nearly every record, so it separated nothing); the number itself stays in its column and `requires_human_review` stays in JSON."""
    if item.get("item_flags"):
        return True
    return any(r != "low-confidence" for r in (item.get("review_reasons") or []))


def _sheet_reasons(item: dict) -> list:
    """The reasons worth showing on the sheet: low confidence is left out (it is true of nearly every row, so it says nothing; the number stays in the Conf. column)."""
    return [r for r in (item.get("review_reasons") or []) if r != "low-confidence"]


def _ref(item: dict) -> str:
    """The citation to show: the paragraph citation (possibly inferred and marked so), else whatever the extractor recorded."""
    return item.get("citation") or item.get("source_ref") or ""


def _section(item: dict) -> str:
    """The section to show: the heading read from the paragraph numbering when there is one (the converter's section path can be nested wrongly), else the section path."""
    return item.get("section_heading") or _join(item.get("section_title_path") or [], " > ")


def _parent_label(item: dict) -> str:
    """"2.17 MAJCOM/DRUs." : the parent paragraph's number and its text, copied from the document; empty when there is none."""
    text = item.get("parent_text") or ""
    return f"{item.get('parent_ref', '')} {text}".strip() if text else ""


def _csv_row(item: dict) -> dict:
    raw = {
        "source_ref": item.get("source_ref", ""),
        "citation": item.get("citation", ""),
        "section_title_path": _join(item.get("section_title_path") or [], " > "),
        "section_heading": item.get("section_heading", ""),
        "page_refs": _join(item.get("page_refs") or [], ", "),
        "applies_to": item.get("applies_to", ""),
        "parent_paragraph": _parent_label(item),
        "source_quote": item.get("source_quote", ""),
        "passage": item.get("passage", ""),
        "audit_question": item.get("audit_question", ""),
        "status": item.get("status", ""),
        "assessor_notes": item.get("assessor_notes", ""),
        "item_flags": _join(item.get("item_flags") or [], "; "),
        "requires_human_review": item.get("requires_human_review", False),
        "review_reasons": _join(item.get("review_reasons") or [], "; "),
        "confidence": item.get("confidence", 0.0),
        "checklist_item_id": item.get("checklist_item_id", ""),
        "requirement_ids": _join(item.get("requirement_ids") or [], ", "),
        "domain_tags": _join(item.get("domain_tags") or [], ", "),
    }
    return {k: _csv_safe(v) for k, v in raw.items()}


def to_csv(checklist: dict) -> str:
    """Return the checklist items as a CSV string (UTF-8, Excel-friendly)."""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=_CSV_COLUMNS, lineterminator="\r\n")
    writer.writeheader()
    for item in checklist.get("items", []) + checklist.get("possible_missed", []):
        writer.writerow(_csv_row(item))
    return buf.getvalue()


def to_json(checklist: dict) -> str:
    """Return the full checklist envelope as a pretty-printed JSON string."""
    return json.dumps(checklist, indent=2, default=str)


def _md_item(item: dict, index: int) -> str:
    section = _section(item)
    pages = _join(item.get("page_refs") or [], ", ")
    header_parts = []
    if section:
        header_parts.append(section)
    if pages:
        header_parts.append(f"p. {pages}")
    heading = " — ".join(header_parts) if header_parts else f"Item {index}"

    lines = [f"## {heading}", ""]

    source_ref = _ref(item)
    if source_ref:
        lines.append(f"**Source Ref:** {source_ref}  ")
    tags = _join(item.get("domain_tags") or [], ", ")
    if tags:
        lines.append(f"**Domain Tags:** {tags}  ")
    conf = item.get("confidence")
    if conf is None:
        conf = 0.0
    lines.append(f"**Confidence:** {conf:.2f}  ")
    lines.append("")

    applies = item.get("applies_to", "")
    if applies:
        lines.append(f"**Applies to:** {applies}  ")
    parent = _parent_label(item)
    if parent:
        lines.append(f"**Parent paragraph:** {parent}  ")
    if applies or parent:
        lines.append("")
    quote = item.get("source_quote", "")
    lines.append(f"> {quote}")
    lines.append("")
    passage = item.get("passage", "")
    if passage:
        lines.append("**Passage** (the requirement is marked >> <<):")
        lines.append("")
        lines.extend(f"    {ln}" for ln in passage.splitlines())
        lines.append("")
    flags = _join(item.get("item_flags") or [], ", ")
    if flags:
        lines.append(f"**Check:** {flags}  ")
        lines.append("")

    audit_q = item.get("audit_question", "")
    lines.append(f"**Audit Question (draft):** {audit_q if audit_q else '*(not generated)*'}  ")

    status = item.get("status", "not-started")
    assessor = item.get("assessor_notes", "")
    lines.append(f"**Status:** {status}  ")
    lines.append(f"**Assessor Notes:** {assessor if assessor else '*(none)*'}  ")
    lines.append("")

    if item.get("requires_human_review"):
        reasons = _join(item.get("review_reasons") or [], "; ")
        lines.append(f"> ⚠ **Requires Review:** {reasons}  ")
        lines.append("")

    req_ids = _join(item.get("requirement_ids") or [], ", ")
    lines.append(f"*ID: {item.get('checklist_item_id', '')} | Req: {req_ids}*")
    lines.append("")
    lines.append("---")
    lines.append("")
    return "\n".join(lines)


# Rows of the workbook: a title and a subtitle, then the column groups and the column headers, then the data
_XLSX_GROUP_ROW = 3
_XLSX_HEADER_ROW = 4
_XLSX_FIRST_DATA_ROW = 5
_STATUS_COL = 9
_CONF_COL = 13


def _xlsx_title(checklist: dict) -> str:
    """The document's name for the sheet: the PDF it was read from, else its id (the same fallback the Markdown export uses)."""
    doc = checklist.get("document") or {}
    return doc.get("source_pdf") or doc.get("document_id") or "Unknown document"


def _when(stamp: object) -> str:
    """"2026-10-11 15:04 UTC" from the envelope's ISO time; the text as it is when it cannot be read."""
    text = str(stamp or "")
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        return text
    suffix = " UTC" if parsed.utcoffset() is not None and parsed.utcoffset().total_seconds() == 0 else ""
    return parsed.strftime("%Y-%m-%d %H:%M") + suffix


def to_xlsx(checklist: dict) -> bytes:
    """Return the checklist items as an Excel workbook (XLSX) in bytes.

    Layout (rows): 1 the document's name, 2 a line of facts (profile, when, how many rows, how many are shaded, how many possible missed), 3 the five column groups (Locate, Ask,
    Record, Verify, Trace), each in its own color, 4 the column headers in the lighter shade of the group's color, then one row per item, then the possible-missed block.
    The columns and their order are the same as before; the Status column is a drop-down and is colored by what is chosen.

    Requires openpyxl (approved WP-23.2 dependency).
    """
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.page import PageMargins

    # Column definitions: (display_header, item_key, col_width, wrap_text)
    _COLS = [
        ("Ref",            "source_ref",            14, False),
        ("Section",        "section_title_path",     26, False),
        ("Pages",          "page_refs",               8, False),
        ("Applies to",     "applies_to",             22, True),
        ("Parent para.",   "parent_text",            30, True),
        ("Requirement",    "source_quote",            44, True),
        ("Passage",        "passage",                 56, True),
        ("Audit Question (draft)", "audit_question",  46, True),
        ("Status",         "status",                  16, False),
        ("Notes",          "assessor_notes",          30, True),
        ("Check",          "item_flags",              20, False),
        ("Reasons",        "review_reasons",          22, False),
        ("Conf.",          "confidence",               7, False),
        ("Item ID",        "checklist_item_id",       26, False),
        ("Req IDs",        "requirement_ids",         20, False),
        ("Tags",           "domain_tags",             20, False),
    ]

    # Column groups: (label, first_col_1based, last_col_1based, band color, header tint). White on the band and dark text on the tint both read at 4.5:1 or better.
    _GROUPS = [
        ("Locate",  1,  5, "1D4ED8", "DBEAFE"),
        ("Ask",     6,  8, "047857", "D1FAE5"),
        ("Record",  9, 10, "B45309", "FEF3C7"),
        ("Verify", 11, 13, "6D28D9", "EDE9FE"),
        ("Trace",  14, 16, "4B5563", "E5E7EB"),
    ]

    face = "Calibri"
    ink, soft = "111827", "4B5563"
    hairline = Side(style="thin", color="D1D5DB")
    box = Border(left=hairline, right=hairline, top=hairline, bottom=hairline)
    _FLAGGED_FILL = PatternFill("solid", fgColor="FFFBEB")

    items_all = checklist.get("items", [])
    missed = checklist.get("possible_missed", [])
    title = _xlsx_title(checklist)

    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = "1D4ED8"
    wb.properties.title = f"Audit checklist: {title}"
    wb.properties.creator = "ReqBot"

    # Rows 1 and 2: which document this is, and the facts about the sheet
    ws.cell(row=1, column=1, value=_csv_safe(f"Audit checklist: {title}")).font = Font(name=face, size=16, bold=True, color=ink)
    ws.row_dimensions[1].height = 26
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
    facts = []
    if checklist.get("profile"):
        facts.append(f"Profile: {checklist['profile']}")
    if checklist.get("generated_at"):
        facts.append(f"Generated: {_when(checklist['generated_at'])}")
    facts.append(f"{len(items_all)} item{'' if len(items_all) == 1 else 's'}")
    shaded = sum(1 for i in items_all if _needs_attention(i))
    if shaded:
        facts.append(f"{shaded} of them shaded for a second look")  # the possible-missed rows below are shaded too: they are counted on their own
    if missed:
        facts.append(f"{len(missed)} possible missed, listed at the bottom")
    ws.cell(row=2, column=1, value=_csv_safe(" · ".join(facts))).font = Font(name=face, size=10, color=soft)
    ws.row_dimensions[2].height = 16

    # Row 3: the column groups, each band in its own color
    for label, c_start, c_end, band, _tint in _GROUPS:
        for col in range(c_start, c_end + 1):
            ws.cell(row=_XLSX_GROUP_ROW, column=col).fill = PatternFill("solid", fgColor=band)
        cell = ws.cell(row=_XLSX_GROUP_ROW, column=c_start, value=label)
        cell.font = Font(name=face, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if c_start != c_end:
            ws.merge_cells(start_row=_XLSX_GROUP_ROW, start_column=c_start, end_row=_XLSX_GROUP_ROW, end_column=c_end)
    ws.row_dimensions[_XLSX_GROUP_ROW].height = 18

    # Row 4: the column headers, in the lighter shade of their group's color
    tint_of = {col: tint for _l, c_start, c_end, _b, tint in _GROUPS for col in range(c_start, c_end + 1)}
    band_of = {col: band for _l, c_start, c_end, band, _t in _GROUPS for col in range(c_start, c_end + 1)}
    for col_idx, (header, _, _, _) in enumerate(_COLS, start=1):
        cell = ws.cell(row=_XLSX_HEADER_ROW, column=col_idx, value=header)
        cell.font = Font(name=face, size=10, bold=True, color=ink)
        cell.fill = PatternFill("solid", fgColor=tint_of[col_idx])
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = Border(left=hairline, right=hairline, top=hairline, bottom=Side(style="medium", color=band_of[col_idx]))
    ws.row_dimensions[_XLSX_HEADER_ROW].height = 30

    # Freeze below the headers, after the Locate group's first three columns (the reference stays in view)
    ws.freeze_panes = f"D{_XLSX_FIRST_DATA_ROW}"

    # Status column data validation (dropdown)
    dv = DataValidation(
        type="list",
        formula1='"not-started,in-progress,compliant,non-compliant,not-applicable"',
        allow_blank=True,
    )
    ws.add_data_validation(dv)

    # Column widths
    for col_idx, (_, _, width, _) in enumerate(_COLS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Tags and confidence are no longer assigned by the pipeline: hide those two columns (kept, so positions do not move) unless some row has a value (older runs do)
    for header, key, present in (("Conf.", "confidence", lambda v: v is not None), ("Tags", "domain_tags", bool)):
        if not any(present(i.get(key)) for i in items_all):
            ws.column_dimensions[get_column_letter(next(n for n, c in enumerate(_COLS, start=1) if c[0] == header))].hidden = True

    # Data rows
    def write_item(item):
        flagged = _needs_attention(item)
        row_fill = _FLAGGED_FILL if flagged else None
        confidence_val = item.get("confidence") or 0.0

        values = [
            _csv_safe(_ref(item)),
            _csv_safe(_section(item)),
            _csv_safe(_join(item.get("page_refs") or [], ", ")),
            _csv_safe(item.get("applies_to") or ""),
            _csv_safe(_parent_label(item)),
            _csv_safe(item.get("source_quote") or ""),
            _csv_safe(item.get("passage") or ""),
            _csv_safe(item.get("audit_question") or ""),
            _csv_safe(item.get("status") or ""),
            _csv_safe(item.get("assessor_notes") or ""),
            _csv_safe(_join(item.get("item_flags") or [], "; ")),
            _csv_safe(_join(_sheet_reasons(item), "; ")),
            confidence_val,
            _csv_safe(item.get("checklist_item_id") or ""),
            _csv_safe(_join(item.get("requirement_ids") or [], ", ")),
            _csv_safe(_join(item.get("domain_tags") or [], ", ")),
        ]

        row_num = ws.max_row + 1
        for col_idx, (_, _, _, wrap) in enumerate(_COLS, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=values[col_idx - 1])
            cell.font = Font(name=face, size=10, bold=(col_idx == 1), color=ink)
            cell.alignment = Alignment(vertical="top", wrap_text=wrap, horizontal="center" if col_idx == _STATUS_COL else None)
            cell.border = box
            if row_fill:
                cell.fill = row_fill

        # Confidence as percentage (value is 0–1 float)
        ws.cell(row=row_num, column=_CONF_COL).number_format = "0%"
        # Status: register with data validation
        dv.add(ws.cell(row=row_num, column=_STATUS_COL))

    for item in items_all:
        write_item(item)

    # Auto-filter covers the extracted rows only; the possible-missed section below is separate
    last_item_row = ws.max_row

    if missed:
        banner_row = ws.max_row + 2
        banner = ws.cell(
            row=banner_row, column=1,
            value=f"POSSIBLE MISSED REQUIREMENTS ({len(missed)}): found in the document text but not extracted. Confirm each; they are not counted above.",
        )
        banner.font = Font(name=face, size=10, bold=True, color="92400E")
        banner.fill = PatternFill("solid", fgColor="FEF3C7")
        banner.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=banner_row, start_column=1, end_row=banner_row, end_column=len(_COLS))
        ws.row_dimensions[banner_row].height = 24
        for item in missed:
            write_item(item)

    # The Status cell takes the color of what was chosen (not-started stays plain)
    status_range = f"{get_column_letter(_STATUS_COL)}{_XLSX_FIRST_DATA_ROW}:{get_column_letter(_STATUS_COL)}{max(ws.max_row, _XLSX_FIRST_DATA_ROW)}"
    for value, fill, ink_color in (
        ("compliant", "DCFCE7", "166534"),
        ("non-compliant", "FEE2E2", "991B1B"),
        ("in-progress", "DBEAFE", "1E40AF"),
        ("not-applicable", "E5E7EB", "4B5563"),
    ):
        ws.conditional_formatting.add(
            status_range,
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill(start_color=fill, end_color=fill, fill_type="solid"), font=Font(color=ink_color, bold=True)),
        )

    # Auto-filter: set after rows are written so range covers all data rows
    ws.auto_filter.ref = f"A{_XLSX_HEADER_ROW}:{get_column_letter(len(_COLS))}{max(last_item_row, _XLSX_HEADER_ROW)}"

    # Printing: landscape letter, the group and header rows repeat on every page, the document's name and the page number sit in the footer
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.print_title_rows = f"{_XLSX_GROUP_ROW}:{_XLSX_HEADER_ROW}"
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6)
    ws.oddFooter.left.text = title[:100].replace("&", "&&")  # an ampersand is a code in a header or footer
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_markdown(checklist: dict) -> str:
    """Return the checklist as a Markdown string."""
    doc = checklist.get("document", {})
    source_pdf = doc.get("source_pdf", "")
    title = source_pdf if source_pdf else doc.get("document_id", "Unknown Document")

    summary = checklist.get("summary", {})
    total = summary.get("total_items", 0)
    review_count = summary.get("items_requiring_review", 0)

    lines = [
        f"# Checklist: {title}",
        "",
        f"**Profile:** {checklist.get('profile', '')}  ",
        f"**Generated:** {checklist.get('generated_at', '')}  ",
        f"**Items:** {total} total, {review_count} requiring review  ",
        "",
        "---",
        "",
    ]

    for i, item in enumerate(checklist.get("items", []), start=1):
        lines.append(_md_item(item, i))

    missed = checklist.get("possible_missed", [])
    if missed:
        lines += [
            "# Possible missed requirements",
            "",
            f"{len(missed)} passage(s) look like obligations but were not extracted (found by a text scan). Confirm each; they are not counted above.",
            "",
            "---",
            "",
        ]
        for i, item in enumerate(missed, start=1):
            lines.append(_md_item(item, i))

    return "\n".join(lines)

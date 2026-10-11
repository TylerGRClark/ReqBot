"""Draw a stage 1 rating sample (rules in README.md) and write the rating sheet. Run `reqbot kinds` for the six documents first.

    PYTHONPATH=. python3 eval/spike_results/requirement_kind/make_rating_sheet.py --seed 61 --exclude outputs/sample.json

The sheet and the sample are named after the rules version and the seed; an existing file is never overwritten. A kind sidecar written by other rules, or one that misses a row, stops the run.
"""
import argparse
import collections
import json
import random
import sys
from pathlib import Path

from services import checklist_service, requirement_kind

HERE = Path(__file__).resolve().parent
PROC = Path.home() / "documents" / "processed"
CHECKLISTS = Path.home() / "reqbot-work" / "checklists"
DOCS = [("AFI 10-2402", "afi10-2402"), ("AFI 13-550", "afi13-550"), ("AFI 17-203", "afi17-203"), ("AFMAN 17-2101", "afman17-2101"), ("AFPD 17-1", "afpd_17-1"),
        ("DAFMAN 17-1305", "dafman17-1305")]
EXTRA = ["duty_with_actor", "duty_imperative", "unknown", "permission"]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=50)
    ap.add_argument("--exclude", action="append", default=[], help="a sample.json whose rows are not drawn again (rows used to revise the rules)")
    args = ap.parse_args()
    version = requirement_kind.RULES_VERSION
    stem = f"{version}_seed{args.seed}"
    sheet, sample_out = CHECKLISTS / f"requirement_kind_rating_{stem}.xlsx", HERE / "outputs" / f"sample_{stem}.json"
    for target in (sheet, sample_out):
        if target.exists():
            sys.exit(f"{target} exists; refusing to overwrite")
    excluded = {r["id"] for f in args.exclude for r in json.loads(Path(f).read_text(encoding="utf-8"))}
    rows = []
    for label, key in DOCS:
        env = checklist_service.generate(PROC, key, "cybersecurity")
        kinds = requirement_kind.load(requirement_kind.sidecar_path(checklist_service.requirements_path(PROC, key)))
        for it in env["items"]:
            if not (it.get("source_quote") or "").strip():
                continue
            rec = kinds.get(it["checklist_item_id"])
            if rec is None or rec.get("rules_version") != version:
                sys.exit(f"{label}: the kind sidecar is missing a row or was written by other rules (needs {version}); run `reqbot kinds --doc {key}` first")
            if it["checklist_item_id"] not in excluded:
                rows.append((label, it, rec))
    by_kind = collections.defaultdict(list)
    for r in rows:
        by_kind[r[2]["kind"]].append(r)
    rng = random.Random(args.seed)
    chosen, short = [], 0
    for kind in requirement_kind.KINDS:
        pool = sorted(by_kind.get(kind, []), key=lambda r: r[1]["checklist_item_id"])
        take = min(2, len(pool))
        short += 2 - take
        chosen += rng.sample(pool, take)
    taken = {r[1]["checklist_item_id"] for r in chosen}
    for kind in (EXTRA * 3)[:max(0, 30 - len(chosen))]:
        pool = [r for r in sorted(by_kind.get(kind, []), key=lambda r: r[1]["checklist_item_id"]) if r[1]["checklist_item_id"] not in taken]
        if pool:
            pick = rng.choice(pool)
            chosen.append(pick)
            taken.add(pick[1]["checklist_item_id"])
    chosen = chosen[:30]
    rng.shuffle(chosen)
    (HERE / "outputs").mkdir(exist_ok=True)
    sample_out.write_text(json.dumps([{"id": r[1]["checklist_item_id"], "doc": r[0], "kind": r[2]["kind"]} for r in chosen], indent=1), encoding="utf-8")
    share = {k: len(v) for k, v in by_kind.items()}
    (HERE / "outputs" / f"counts_{version}.json").write_text(json.dumps({"rows": len(rows), "excluded": len(excluded), "by_kind": share}, indent=1), encoding="utf-8")
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rate"
    ws.append(["Document", "Ref", "Requirement", "Kind given", "Why (the rule that fired)", "Tags", "Right / Wrong / Unclear", "If wrong, the kind you would call it", "Surrounding text", "Item ID"])
    for label, it, rec in chosen:
        ws.append([label, it.get("citation", ""), it["source_quote"], rec["kind"], "; ".join(rec["kind_reasons"]), ", ".join(rec["kind_tags"]), "", "", it.get("passage", ""), it["checklist_item_id"]])
    for i, w in enumerate([14, 12, 62, 22, 40, 18, 16, 26, 70, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    wb.save(sheet)
    print(f"rows in the population: {len(rows)}; shares: {share}; sample: {len(chosen)} (places a kind could not fill: {short}); sheet: {sheet}")


if __name__ == "__main__":
    sys.exit(main())

"""Draft audit question prompt v4 experiment. Rules and pass criteria are in README.md (fixed before any v4 output was generated).

Run: PYTHONPATH=. python3 eval/spike_results/question_prompt_v4/run_prompt_v4.py --ollama-url http://192.168.90.100:11434
Resumable: rows already in outputs/runs.jsonl are not called again.
"""
import argparse
import json
import random
import re
import sys
from pathlib import Path

from services import audit_questions as aq
from services import checklist_service as cs

HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"
PROC = Path.home() / "documents" / "processed"
DOCS = [("AFI 10-2402", "afi10-2402"), ("AFI 13-550", "afi13-550"), ("AFI 17-203", "afi17-203")]
SHEET = Path.home() / "reqbot-work" / "checklists" / "question_prompt_v4_rating.xlsx"

SET_A = ["MISS-7f3e5f9ff07af443", "MISS-577ceab6af9f4608", "MISS-d7b75d617c5b47b2", "CHK-689caa8b6d5e1d3c", "MISS-ed6db2bcccae56ee", "MISS-dc86fa45974f9d78",
         "MISS-3952678e90428fae", "MISS-70dd20b87b6720f7", "MISS-f467a1e85de27479", "MISS-8644c47719a51b64", "CHK-443cb500a3a14018", "MISS-e610956e5a8e6abc",
         "MISS-5470ef430c84b16c", "MISS-2716f46eaab91b7b", "MISS-420dcc7062abac3a", "CHK-c142ed10711872a6", "MISS-90bbcd60e5ef5090", "CHK-68601159c0f363b1",
         "CHK-41e0cd12bd7a677b", "CHK-ce5ab5b4f86022be", "MISS-ec5a2071dcf63270"]  # rows 2 to 22 of afi_all_flagged_rows.xlsx, in sheet order
TARGET_ID = "MISS-ec5a2071dcf63270"  # sheet row 22: "Required courses will be designated by the HAF CARM Program office." (2.17.26)

RULE2_V3 = ('- Name who must act if the material says so (the "Applies to" line or the requirement itself). If the requirement does not say who acts and nothing below does '
            'either, ask about "the unit".\n')
RULE2_V4 = ('- Name who must act. If the requirement itself says who acts (for example "by the program office" or "the commander will"), use that; the "Applies to" line is the '
            'actor only when the requirement names none. If neither says, ask about "the unit".\n')
SCOPE_V4 = ("- Ask only about the requirement marked >> << in the passage. The sentences around it are background: do not ask about them and do not put their duties into "
            "the question.\n")
assert RULE2_V3 in aq.PROMPT
PROMPT_V4 = aq.PROMPT.replace(RULE2_V3, SCOPE_V4 + RULE2_V4)


def build_prompt_v4(item: dict) -> str:
    """aq.build_prompt with the v4 template: the same inputs, the same prohibition rule."""
    template = PROMPT_V4
    if aq.PROHIBITION.search(item.get("source_quote") or ""):
        template = template.replace("- If the requirement is a fragment,", aq.PROHIBITION_RULE + "- If the requirement is a fragment,")
    parent = f"{item.get('parent_ref') or ''} {item.get('parent_text') or ''}".strip()
    return template.format(applies=item.get("applies_to") or "(none)", parent=parent or "(none)", quote=item.get("source_quote") or "",
                           passage=(item.get("passage") or "")[:1500])


def before_text(item: dict) -> str:
    """The text on the passage line that holds the marked sentence, ahead of the marker, without the paragraph number."""
    line = next((ln for ln in (item.get("passage") or "").splitlines() if ">>" in ln), "")
    before = line.split(">>")[0]
    return re.sub(r"^\W*(?:[A-Z]{0,2}\d+(?:\.\d+)*\.?\s*)?", "", before.strip()).strip()


def eligible(item: dict) -> bool:
    return bool((item.get("source_quote") or "").strip()) and not aq.skipped_because(item)


_STOP = {"which", "their", "there", "would", "should", "shall", "about", "other", "these", "those", "while", "where", "being", "under", "after", "before", "within", "upon",
         "whether", "unit", "yes"}


def _stems(text: str) -> set:
    return {re.sub(r"(ing|ed|es|s)$", "", w) for w in (m.lower() for m in re.findall(r"[A-Za-z]{5,}", text or "")) if w not in _STOP}


def blend_words(item: dict, question: str) -> list:
    """Words of the question found in the sentence before the marked one and nowhere in the marked sentence, its "Applies to" line or its parent paragraph."""
    own = _stems(" ".join([item.get("source_quote") or "", item.get("applies_to") or "", item.get("parent_text") or ""]))
    return sorted((_stems(question) & _stems(before_text(item))) - own)


def norm(q: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", (q or "").lower())).strip()


def ask(prompt: str, url: str):
    for _ in range(2):
        try:
            return aq.parse_answer(aq.ollama_call(prompt, aq.MODEL, url)), None
        except (ValueError, KeyError) as e:
            err = f"malformed response: {e}"
        except OSError as e:
            return None, f"model call failed: {e}"
    return None, err


def pick_rows():
    rows = []  # (set, doc label, item)
    allitems = []
    for label, key in DOCS:
        for it in cs.generate(PROC, key, "cybersecurity")["items"]:
            allitems.append((label, it))
    by_id = {it["checklist_item_id"]: (label, it) for label, it in allitems}
    a_ids = set(SET_A)
    for cid in SET_A:
        label, it = by_id[cid]
        rows.append(("A", label, it))
    pool = [(label, it) for label, it in allitems if it["checklist_item_id"] not in a_ids and eligible(it)]
    b = [(label, it) for label, it in pool if len(before_text(it)) >= 25]
    b_ids = {it["checklist_item_id"] for _, it in b}
    rest = sorted([(label, it) for label, it in pool if it["checklist_item_id"] not in b_ids], key=lambda t: (t[0], t[1]["checklist_item_id"]))
    c = random.Random(47).sample(rest, 40)
    rows += [("B", label, it) for label, it in b] + [("C", label, it) for label, it in c]
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ollama-url", required=True)
    args = ap.parse_args()
    OUT.mkdir(exist_ok=True)
    runs_path = OUT / "runs.jsonl"
    done = {}
    if runs_path.exists():
        for line in runs_path.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            done[r["id"]] = r
    rows = pick_rows()
    print(f"rows: A={sum(1 for s, _, _ in rows if s == 'A')} B={sum(1 for s, _, _ in rows if s == 'B')} C={sum(1 for s, _, _ in rows if s == 'C')}", flush=True)
    for n, (s, label, it) in enumerate(rows, 1):
        cid = it["checklist_item_id"]
        rec = done.get(cid) or {"id": cid, "set": s, "doc": label, "ref": it.get("citation", ""), "quote": it["source_quote"], "v3_stored": it.get("audit_question") or None,
                                "gate": aq.skipped_because(it)}
        if not eligible(it):
            rec["note"] = "withheld by the code gate: not compared"
        else:
            if "v4" not in rec and "v4_error" not in rec:
                rec["v4"], err = ask(build_prompt_v4(it), args.ollama_url)
                if err:
                    rec["v4_error"] = err
            if s in ("A", "C") and "v3_rerun" not in rec and "v3_rerun_error" not in rec:
                rec["v3_rerun"], err = ask(aq.build_prompt(it), args.ollama_url)
                if err:
                    rec["v3_rerun_error"] = err
            rec["blend_v3"] = blend_words(it, rec["v3_stored"] or "")
            rec["blend_v4"] = blend_words(it, rec.get("v4") or "")
            rec["unver_v3"] = aq.unverified_terms(rec["v3_stored"], aq.row_material(it)) if rec["v3_stored"] else []
            rec["unver_v4"] = aq.unverified_terms(rec.get("v4"), aq.row_material(it)) if rec.get("v4") else []
        rec["passage"] = it.get("passage", "")
        done[cid] = rec
        if n % 10 == 0 or n == len(rows):
            runs_path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in done.values()), encoding="utf-8")
            print(f"  {n}/{len(rows)}", flush=True)
    runs_path.write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in done.values()), encoding="utf-8")
    measures(list(done.values()))
    rating_sheet(list(done.values()))


def measures(recs):
    out = {}
    for s in ("A", "B", "C", "all"):
        rs = [r for r in recs if (s == "all" or r["set"] == s) and "note" not in r and ("v4" in r or "v4_error" in r)]
        if not rs:
            continue
        ok = [r for r in rs if "v4_error" not in r]
        d = {"rows": len(rs), "v4_errors": len(rs) - len(ok),
             "changed_v4_vs_v3": sum(norm(r["v4"]) != norm(r["v3_stored"]) for r in ok),
             "null_v3": sum(1 for r in ok if not r["v3_stored"]), "null_v4": sum(1 for r in ok if not r["v4"]),
             "unverified_v3": sum(1 for r in ok if r["unver_v3"]), "unverified_v4": sum(1 for r in ok if r["unver_v4"]),
             "blend_suspect_v3": sum(1 for r in ok if len(r["blend_v3"]) >= 2), "blend_suspect_v4": sum(1 for r in ok if len(r["blend_v4"]) >= 2)}
        rr = [r for r in ok if "v3_rerun" in r]
        if rr:
            d["noise_floor_rows"] = len(rr)
            d["noise_floor_changed_v3_rerun_vs_v3"] = sum(norm(r["v3_rerun"]) != norm(r["v3_stored"]) for r in rr)
            d["v4_changed_on_same_rows"] = sum(norm(r["v4"]) != norm(r["v3_stored"]) for r in rr)
        out[s] = d
    (OUT / "measures.json").write_text(json.dumps(out, indent=2), encoding="utf-8")
    print(json.dumps(out, indent=2))


def rating_sheet(recs):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    changed = {s: sorted([r for r in recs if r["set"] == s and "note" not in r and "v4" in r and norm(r["v4"]) != norm(r["v3_stored"])], key=lambda r: r["id"]) for s in "ABC"}
    random.Random(48).shuffle(changed["B"])
    random.Random(48).shuffle(changed["C"])
    chosen = [r for r in recs if r["set"] == "A" and r in changed["A"]] + changed["B"][:15] + changed["C"][:10]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rate"
    head = ["Set", "Document", "Ref", "Requirement (marked sentence)", "Question now (v3)", "New question (v4)", "Better / Same / Worse", "Notes", "Passage", "Item ID"]
    ws.append(head)
    for r in chosen:
        ws.append([r["set"], r["doc"], r["ref"], r["quote"], r["v3_stored"] or "(none)", r["v4"] or "(none)", "", "", r["passage"], r["id"]])
    for i, w in enumerate([5, 13, 12, 55, 50, 50, 14, 28, 80, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"
    wb.save(SHEET)
    print(f"rating sheet: {len(chosen)} rows -> {SHEET}")


if __name__ == "__main__":
    sys.exit(main())

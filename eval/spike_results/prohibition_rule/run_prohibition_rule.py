"""Prohibition rule experiment. Rules are in README.md (fixed before any new question was generated).

Run: PYTHONPATH=. python3 eval/spike_results/prohibition_rule/run_prohibition_rule.py --ollama-url http://192.168.90.100:11434
"""
import argparse
import json
import re
import sys
from pathlib import Path

from services import audit_questions as aq
from services import checklist_service as cs

HERE = Path(__file__).resolve().parent
OUT = HERE / "outputs"
PROC = Path.home() / "documents" / "processed"
SHEET = Path.home() / "reqbot-work" / "checklists" / "prohibition_rule_rating.xlsx"
AFIS = ["afi10-2402", "afi13-550", "afi17-203"]
LABEL = {"afi10-2402": "AFI 10-2402", "afi13-550": "AFI 13-550", "afi17-203": "AFI 17-203"}

NEW = re.compile(
    r"\b(?:shall|must|should|will|may|can|could|would)\s+not\b(?!\s+(?:limited|only|necessarily))|\bcannot\b|\bnever\b|\bprohibited\b|\bforbidden\b"
    r"|\b(?:is|are)\s+not\s+(?:authorized|permitted|allowed|to)\b|^\W*(?:\(\w+\)\s*)?(?:do|does)\s+not\b", re.IGNORECASE)
AVOID = re.compile(r"\bavoid(?:s|ing)?\b", re.IGNORECASE)


def prompt_with(item: dict, fires: bool) -> str:
    """aq.build_prompt with the prohibition rule inserted or not, as decided by `fires`."""
    template = aq.PROMPT
    if fires:
        template = template.replace("- If the requirement is a fragment,", aq.PROHIBITION_RULE + "- If the requirement is a fragment,")
    parent = f"{item.get('parent_ref') or ''} {item.get('parent_text') or ''}".strip()
    return template.format(applies=item.get("applies_to") or "(none)", parent=parent or "(none)", quote=item.get("source_quote") or "", passage=(item.get("passage") or "")[:1500])


def norm(q):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", (q or "").lower())).strip()


def ask(prompt, url):
    for _ in range(2):
        try:
            return aq.parse_answer(aq.ollama_call(prompt, aq.MODEL, url)), None
        except (ValueError, KeyError) as e:
            err = f"malformed response: {e}"
        except OSError as e:
            return None, f"model call failed: {e}"
    return None, err


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ollama-url", required=True)
    args = ap.parse_args()
    OUT.mkdir(exist_ok=True)
    keys = sorted({p.name.rsplit("_2026", 1)[0] for p in PROC.iterdir() if p.is_dir() and "_2026" in p.name} - {"CNSSI_No1253"})
    released, same_prompt = [], 0
    for key in keys:
        for it in cs.generate(PROC, key, "cybersecurity")["items"]:
            q = it.get("source_quote") or ""
            if not q.strip() or aq.skipped_because(it):
                continue
            old, new = bool(aq.PROHIBITION.search(q)), bool(NEW.search(q))
            assert not (new and not old), "the narrowed pattern must be a subset of the old one"
            if old and not new:
                released.append((key, it))
            else:
                assert prompt_with(it, new) == aq.build_prompt(it), "a row that keeps its rule state must keep a byte-identical prompt"
                same_prompt += 1
    print(f"released rows: {len(released)}; rows with an identical prompt (asserted): {same_prompt}", flush=True)
    recs = []
    for n, (key, it) in enumerate(released, 1):
        cur = it.get("audit_question") or None
        new_q, err = ask(prompt_with(it, False), args.ollama_url)
        rerun, err2 = ask(prompt_with(it, True), args.ollama_url)
        recs.append({"id": it["checklist_item_id"], "doc": key, "ref": it.get("citation", ""), "quote": it["source_quote"], "applies_to": it.get("applies_to", ""),
                     "current": cur, "new": new_q, "rerun_current_prompt": rerun, "errors": [e for e in (err, err2) if e], "passage": it.get("passage", "")})
        if n % 10 == 0 or n == len(released):
            print(f"  {n}/{len(released)}", flush=True)
    (OUT / "runs.jsonl").write_text("".join(json.dumps(r, ensure_ascii=False) + "\n" for r in recs), encoding="utf-8")
    measures(recs, same_prompt)
    rating_sheet(recs)


def measures(recs, same_prompt):
    ok = [r for r in recs if not r["errors"]]
    isnull = lambda q: aq.no_question(q)  # noqa: E731
    m = {"released_rows": len(recs), "errors": len(recs) - len(ok), "identical_prompt_rows_asserted": same_prompt,
         "avoid_current": sum(1 for r in ok if r["current"] and AVOID.search(r["current"])), "avoid_new": sum(1 for r in ok if r["new"] and AVOID.search(r["new"])),
         "null_current": sum(1 for r in ok if isnull(r["current"])), "null_new": sum(1 for r in ok if isnull(r["new"])),
         "changed_new_vs_current": sum(1 for r in ok if norm(r["new"]) != norm(r["current"])),
         "noise_floor_changed_rerun_vs_current": sum(1 for r in ok if norm(r["rerun_current_prompt"]) != norm(r["current"]))}
    (OUT / "measures.json").write_text(json.dumps(m, indent=2), encoding="utf-8")
    print(json.dumps(m, indent=2))


def rating_sheet(recs):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ordered = [r for r in recs if r["doc"] in AFIS] + [r for r in recs if r["doc"] not in AFIS]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rate"
    ws.append(["Document", "Ref", "Requirement", "Question now", "New question", "Better / Same / Worse", "Notes", "Passage", "Item ID"])
    for r in ordered:
        ws.append([LABEL.get(r["doc"], r["doc"]), r["ref"], r["quote"], r["current"] or "(none)", r["new"] or "(none)", "", "", r["passage"], r["id"]])
    for i, w in enumerate([14, 12, 60, 50, 50, 14, 28, 80, 22], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    wb.save(SHEET)
    print(f"rating sheet: {len(ordered)} rows -> {SHEET}")


if __name__ == "__main__":
    sys.exit(main())

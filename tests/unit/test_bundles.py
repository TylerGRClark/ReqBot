import json
from pathlib import Path

import openpyxl
import pytest

from pipeline import bundle_export
from services import bundles


def _write(path: Path, data: dict) -> Path:
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def test_the_users_file_wins_by_name_and_adds_new_bundles(tmp_path):
    starters = {"a": {"label": "Starter A", "documents": [{"doc": "x"}]}, "b": {"label": "B", "documents": ["y"]}}
    user = _write(tmp_path / "user.json", {"bundles": {"a": {"label": "My A", "documents": [{"doc": "z", "title": "Z"}]}, "c": {"label": "C", "documents": []}}})
    table = bundles.load_bundles(user, starters)
    assert sorted(table) == ["a", "b", "c"]
    assert table["a"]["label"] == "My A" and table["a"]["documents"] == [{"doc": "z", "title": "Z"}]
    assert table["b"]["documents"] == [{"doc": "y", "title": ""}]


def test_a_missing_or_broken_file_is_an_empty_list_not_an_error(tmp_path):
    (tmp_path / "bad.json").write_text("{not json", encoding="utf-8")
    assert bundles.load_bundles(tmp_path / "bad.json", {}) == {}
    assert bundles.load_bundles(tmp_path / "nope.json", {}) == {}


def test_add_and_remove_by_document_key_ignoring_case_and_duplicates():
    table = {"u": {"label": "U", "note": "", "documents": [{"doc": "afi17-101", "title": "T"}, {"doc": "afi17-130", "title": ""}]}}
    got = bundles.resolve("u", add=["DODI 5200.48", "AFI17-101"], remove=["AFI17-130"], bundles=table)
    assert [d["doc"] for d in got["documents"]] == ["afi17-101", "DODI 5200.48"]
    assert table["u"]["documents"][1]["doc"] == "afi17-130"  # the stored bundle is not changed


def test_an_unknown_bundle_name_is_a_key_error():
    with pytest.raises(KeyError):
        bundles.resolve("nope", bundles={})


def test_save_keeps_the_users_other_bundles_and_rejects_a_bad_name(tmp_path):
    path = tmp_path / "u.json"
    bundles.save_user_bundle("first", "First", ["a"], path=path)
    bundles.save_user_bundle("second", "Second", ["b", "c"], path=path)
    saved = json.loads(path.read_text())["bundles"]
    assert sorted(saved) == ["first", "second"] and [d["doc"] for d in saved["second"]["documents"]] == ["b", "c"]
    with pytest.raises(ValueError):
        bundles.save_user_bundle("Bad Name!", "x", [], path=path)


def test_the_shipped_starter_names_only_publications_with_a_title():
    starter = bundles.load_bundles(Path("/nonexistent/bundles.json"))
    assert starter, "the starter file must define at least one bundle"
    for b in starter.values():
        assert b["label"] and b["documents"]
        assert all(d["doc"] and d["title"] for d in b["documents"])


def test_workbook_has_a_contents_tab_and_a_tab_per_loaded_publication(tmp_path, monkeypatch):
    envelopes = {
        "doc-a": {"document": {"source_pdf": "doc-a.pdf"}, "items": [
            {"source_quote": "Do A.", "audit_question": "Is A done?", "checklist_item_id": "CHK-1", "requirement_ids": ["R1"]},
            {"source_quote": "Do B.", "audit_question": "", "checklist_item_id": "CHK-2", "requirement_ids": ["R2"]}]},
    }

    def fake_generate(processed_dir, key, profile):
        if key not in envelopes:
            raise FileNotFoundError(key)
        return envelopes[key]

    monkeypatch.setattr(bundle_export.checklist_service, "generate", fake_generate)
    bundle = {"label": "Test bundle", "note": "n", "documents": [{"doc": "doc-a", "title": "Doc A, the first"}, {"doc": "doc-b", "title": "Doc B"}]}
    data, summary = bundle_export.build_bundle_workbook(tmp_path, bundle)
    assert [(r["doc"], r["loaded"], r["rows"], r["with_question"]) for r in summary] == [("doc-a", True, 2, 1), ("doc-b", False, 0, 0)]
    path = tmp_path / "out.xlsx"
    path.write_bytes(data)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Contents", "doc-a"]
    rows = [[c for c in r] for r in wb["Contents"].iter_rows(min_row=5, max_row=6, values_only=True)]
    assert rows[0][:5] == ["doc-a", "Doc A, the first", 2, 1, "yes"]
    assert rows[1][:5] == [None, "Doc B", None, None, "not loaded yet"]
    assert wb["doc-a"].cell(row=4, column=2).value == "Requirement"


def test_tab_titles_are_unique_short_and_safe():
    used: set = set()
    a = bundle_export._tab_title("DODI 5200.48", used)
    b = bundle_export._tab_title("dodi 5200.48", used)
    c = bundle_export._tab_title("x" * 50 + "/[y]", used)
    assert a == "DODI 5200.48" and b == "dodi 5200.48_2" and len(c) <= 31 and not any(ch in c for ch in "[]:*?/\\")


def test_an_unknown_profile_is_an_error_not_a_list_of_unloaded_publications(tmp_path):
    with pytest.raises(FileNotFoundError):
        bundle_export.build_bundle_workbook(tmp_path, {"label": "L", "documents": [{"doc": "doc-a", "title": "A"}]}, profile="no-such-profile")


def test_a_label_note_or_title_that_looks_like_a_formula_is_written_as_text(tmp_path, monkeypatch):
    def fake_generate(processed_dir, key, profile):
        raise FileNotFoundError(key)

    monkeypatch.setattr(bundle_export.checklist_service, "generate", fake_generate)
    bundle = {"label": "=HYPERLINK(\"http://evil.example\")", "note": "+cmd|' /C calc'!A0", "documents": [{"doc": "doc-a", "title": "@SUM(1+1)"}]}
    data, _ = bundle_export.build_bundle_workbook(tmp_path, bundle)
    path = tmp_path / "f.xlsx"
    path.write_bytes(data)
    ws = openpyxl.load_workbook(path)["Contents"]
    for cell in (ws["A1"], ws["A2"], ws["B5"]):
        assert not str(cell.value).startswith(("=", "+", "-", "@")), cell.value

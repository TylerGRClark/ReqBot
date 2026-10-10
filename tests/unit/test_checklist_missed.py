"""WP-46.2: the rule-based scan for passages that look like obligations but were not extracted."""
import csv
import io
import json

import openpyxl

from pipeline.checklist_export import to_csv, to_markdown, to_xlsx
from services import checklist_missed as M
from services.checklist_service import generate

CHUNKS = {
    1: {"chunk_id": 1, "page_start": 4, "section_title_path": ["ROLES AND RESPONSIBILITIES", "2.3. AF/A3"], "raw_text": (
        "2.3.1. AF/A3 will:\n"
        "2.3.2. The CFP will notify the servicing MCCC within 24 hours of detection.\n"
        "2.3.3. Reports are filed in the system of record.\n"
        "2.3.4. Support the ACD unit as directed during investigation of the incident.\n"
        "2.3.5. No event will be closed out as a Category 8 without review by the NOS.")},
    2: {"chunk_id": 2, "page_start": 5, "section_title_path": ["Intro"], "raw_text": "This section describes the process. The CFP will notify the servicing MCCC within 24 hours of detection."},
    3: {"chunk_id": 3, "page_start": 6, "section_title_path": ["Waivers"], "raw_text": "Submit requests for waivers through the chain of command to the appropriate tier."},
}


def test_units_split_at_paragraph_markers_else_at_sentence_ends():
    assert len(M.paragraph_units(CHUNKS[1]["raw_text"])) == 5
    assert M.paragraph_units("One sentence here. Another one follows; and a third.")[0] == "One sentence here."


def test_only_uncovered_obligation_looking_paragraphs_are_listed_once():
    quotes = ["The CFP will notify the servicing MCCC within 24 hours of detection."]
    found = M.find_possible_missed(CHUNKS, quotes)
    texts = [f["source_quote"] for f in found]
    assert any(t.startswith("2.3.4. Support the ACD unit") for t in texts)  # imperative, not covered
    assert any("No event will be closed out" in t for t in texts)  # modal, not covered
    assert not any("AF/A3 will:" in t for t in texts)  # a lead-in (ends with a colon) is carried by its items
    assert not any("Reports are filed" in t for t in texts)  # no modal, no imperative opener
    assert not any("CFP will notify" in t for t in texts)  # covered by an extracted quote; the repeat in chunk 2 is covered too
    first = next(f for f in found if "Support the ACD" in f["source_quote"])
    assert first["source_ref"] == "2.3.4" and first["page_refs"] == [4] and first["applies_to"] == "AF/A3"
    assert first["item_flags"] == ["possible_missed", "imperative"] and first["checklist_item_id"].startswith("MISS-") and first["requirement_ids"] == []


def test_coverage_ignores_spacing_before_punctuation():
    chunks = {1: {"chunk_id": 1, "page_start": 1, "section_title_path": [], "raw_text": "The CFP shall review the access logs ( weekly ) and report findings ."}}
    assert M.find_possible_missed(chunks, ["The CFP shall review the access logs (weekly) and report findings."]) == []
    assert len(M.find_possible_missed(chunks, [])) == 1


def _write(tmp_path, quotes, chunks):
    run_dir = tmp_path / "doc_20260101_120000"
    run_dir.mkdir()
    recs = [{"requirement_id": f"REQ-{i}", "source_quote": q, "source_ref": "2.3", "chunk_id": 1, "section_title_path": ["ROLES"], "domain_tags": ["x"], "confidence": 0.9,
             "page_start": 4, "page_end": 4} for i, q in enumerate(quotes)]
    (run_dir / "doc_requirements_normalized.jsonl").write_text("".join(json.dumps(r) + "\n" for r in recs))
    if chunks:
        (run_dir / "doc_chunks.jsonl").write_text("".join(json.dumps(c) + "\n" for c in chunks.values()))


def test_numbered_duty_sentences_become_rows_in_document_order_and_the_rest_stay_apart(tmp_path):
    _write(tmp_path, ["The CFP will notify the servicing MCCC within 24 hours of detection."], CHUNKS)
    result = generate(tmp_path, "doc", "cybersecurity")
    items = result["items"]
    found = [i for i in items if "found_by_text_scan" in i["item_flags"]]
    assert [i["citation"] for i in items] == ["2.3", "2.3.4", "2.3.5"]  # the extracted row first, then the two found rows after it in the chunk
    assert [i["source_quote"] for i in found] == ["Support the ACD unit as directed during investigation of the incident.", "No event will be closed out as a Category 8 without review by the NOS."]
    row = found[0]
    assert row["item_flags"][0] == "found_by_text_scan" and row["requirement_ids"] == [] and row["extracted_quote"] == "" and row["review_reasons"] == ["not-extracted"]
    assert row["checklist_item_id"].startswith("MISS-") and row["page_refs"] == [4] and "AF/A3" in row["applies_to"] and ">> Support the ACD unit" in row["passage"]
    assert result["summary"]["total_items"] == 3 and result["summary"]["found_by_text_scan"] == 2
    assert [m["source_quote"] for m in result["possible_missed"]] == ["Submit requests for waivers through the chain of command to the appropriate tier."]  # no paragraph number: stays in the block
    assert result["summary"]["possible_missed"] == 1 and all("possible_missed" in m["item_flags"] for m in result["possible_missed"])


def test_only_the_sentence_that_holds_the_duty_is_promoted_not_the_paragraph():
    raw = ("3.7. Response and Recovery. Response and recovery includes the detailed steps taken to restore the integrity of affected systems. "
           "The local CFP, with the assistance of the mission owner, will develop a Plan of Action and Milestones (POA&M) detailing the required actions (T-2) .\n"
           "3.8. Training. Personnel are trained every year in the use of the reporting tools.")
    promoted, rest = M.split_candidates({1: {"chunk_id": 1, "page_start": 2, "section_title_path": ["Actions"], "raw_text": raw}}, [])
    assert rest == [] and len(promoted) == 1
    assert promoted[0]["text"] == "The local CFP, with the assistance of the mission owner, will develop a Plan of Action and Milestones (POA&M) detailing the required actions (T-2)."
    assert promoted[0]["ref"] == "3.7" and raw[promoted[0]["offset"]:].startswith("The local CFP")


def test_an_opening_include_is_a_duty_but_a_purpose_sentence_is_not():
    raw = ("3.6. Incident Analysis. Incident analysis is a series of analytical steps taken to find out what happened. Include the mission owner in the process. "
           "The purpose of this analysis is to understand the technical details of the incident.\n3.7. Training. Personnel are trained every year.")
    promoted, rest = M.split_candidates({1: {"chunk_id": 1, "page_start": 1, "section_title_path": [], "raw_text": raw}}, [])
    assert [p["text"] for p in promoted] == ["Include the mission owner in the process."] and rest == []
    # "includes ..." is a description, and a paragraph with no duty sentence of its own stays in the block rather than disappearing
    only_purpose = "3.6. Incident Analysis. The purpose of this analysis is to understand the technical details of the incident.\n3.7. Training. Personnel are trained every year."
    promoted, rest = M.split_candidates({1: {"chunk_id": 1, "page_start": 1, "section_title_path": [], "raw_text": only_purpose}}, [])
    assert promoted == [] and len(rest) == 1


def test_a_paragraph_whose_duty_is_already_extracted_with_a_lead_in_is_not_listed_at_all():
    raw = "2.2.12.2.  Execute Lead Command oversight and management of the AN/USQ-225. (T-1)"
    chunks = {1: {"chunk_id": 1, "page_start": 4, "section_title_path": [], "raw_text": raw}}
    extracted = ["AFGSC will: Execute Lead Command oversight and management of the AN/USQ-225."]  # the extracted row has a lead-in attached and no tier tag
    assert len(M.find_possible_missed(chunks, extracted)) == 1  # the paragraph-level scan could not tell
    assert M.split_candidates(chunks, extracted) == ([], [])
    assert len(M.split_candidates(chunks, [])[0]) == 1  # not extracted: found


def test_a_tier_tag_before_the_full_stop_does_not_hide_an_extracted_sentence():
    raw = "3.4. Notification messages must be properly classified (see paragraph 3.4.4.3 below) (T-2) .\n3.5. Training. Personnel are trained every year."
    chunks = {1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}
    extracted = ["Notification messages must be properly classified (see paragraph 3.4.4.3 below)."]  # the extracted row stops before the tag
    assert M.normalize("units (T-2) .") == M.normalize("units.") == "units."
    assert M.find_possible_missed(chunks, extracted) == [] and M.split_candidates(chunks, extracted) == ([], [])
    assert len(M.split_candidates(chunks, [])[0]) == 1


def test_an_extracted_sentence_does_not_hide_the_other_duty_sentences_of_its_paragraph():
    raw = ("3.5.1. The CFP will isolate the affected system within one hour of detection. The MCCC must notify the 624 OC of the isolation and its cause.\n"
           "3.5.2. Training. Personnel are trained every year.")
    chunks = {1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}
    extracted = ["The CFP will isolate the affected system within one hour of detection."]  # one of the two duty sentences is extracted; it is contained in the paragraph
    assert M.find_possible_missed(chunks, extracted) == []  # the paragraph-level scan stops there
    promoted, rest = M.split_candidates(chunks, extracted)
    assert [p["text"] for p in promoted] == ["The MCCC must notify the 624 OC of the isolation and its cause."] and rest == []


def test_identical_wording_in_two_paragraphs_gives_two_rows_with_their_own_position_and_id():
    raw = ("2.6.2. Provide a representative to the working group when requested.\n2.8.2. Provide a representative to the working group when requested.\n"
           "2.9. Training. Personnel are trained every year.")
    promoted, _ = M.split_candidates({1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}, [])
    assert [p["ref"] for p in promoted] == ["2.6.2", "2.8.2"] and promoted[0]["text"] == promoted[1]["text"]
    assert promoted[0]["offset"] < promoted[1]["offset"] and raw[:promoted[1]["offset"]].rstrip().endswith("2.8.2.")  # each is placed at its own paragraph, the second after its own number
    assert promoted[0]["checklist_item_id"] != promoted[1]["checklist_item_id"]


def test_a_later_imperative_sentence_qualifies_a_paragraph_that_has_no_modal():
    raw = "3.6. Incident Analysis. Incident analysis describes the process of finding out what happened. Include the mission owner.\n3.7. Training. Personnel are trained every year."
    chunks = {1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}
    assert M.find_possible_missed(chunks, []) == []  # the paragraph-level scan looks at the first word only
    promoted, _ = M.split_candidates(chunks, [])
    assert [p["text"] for p in promoted] == ["Include the mission owner."] and promoted[0]["ref"] == "3.6"


def test_the_sentences_of_a_chunk_with_one_numbered_paragraph_keep_its_number():
    raw = "3.6. Incident Analysis. Incident analysis describes the process. The CFP will notify the MCCC of every incident. Include the mission owner in the process."
    promoted, rest = M.split_candidates({1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}, [])
    assert [(p["ref"], p["text"]) for p in promoted] == [("3.6", "The CFP will notify the MCCC of every incident."), ("3.6", "Include the mission owner in the process.")] and rest == []
    no_number = "Incident analysis describes the process. The CFP will notify the MCCC of every incident."
    promoted, rest = M.split_candidates({1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": no_number}}, [])
    assert promoted == [] and len(rest) == 1  # no number anywhere in the chunk: nothing to inherit, so it stays in the block


def test_a_found_row_is_marked_in_its_own_paragraph_when_the_wording_repeats(tmp_path):
    run_dir = tmp_path / "doc_20260101_120000"
    run_dir.mkdir()
    raw = "2.6.2. Provide a representative to the working group when requested.\n2.8.2. Provide a representative to the working group when requested.\n2.9. Training. Personnel are trained every year."
    (run_dir / "doc_requirements_normalized.jsonl").write_text(json.dumps({"requirement_id": "REQ-1", "source_quote": "Personnel are trained every year.", "source_ref": "2.9", "chunk_id": 1,
                                                                            "section_title_path": [], "domain_tags": [], "confidence": None, "page_start": 2, "page_end": 2}) + "\n")
    (run_dir / "doc_chunks.jsonl").write_text(json.dumps({"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}) + "\n")
    items = generate(tmp_path, "doc", "cybersecurity")["items"]
    first, second = [i for i in items if "found_by_text_scan" in i["item_flags"]]
    assert first["citation"] == "2.6.2" and second["citation"] == "2.8.2"
    assert first["passage"].count(">>") == 1 and second["passage"].count(">>") == 1
    assert first["passage"].index(">>") < first["passage"].index("2.8.2.") and second["passage"].index(">>") > second["passage"].index("2.8.2.")  # the second row's mark is after its own number


def test_a_long_duty_sentence_is_promoted_whole_and_can_be_located(tmp_path):
    long_sentence = "The local CFP will " + ", ".join(f"coordinate step {i} of the response with the mission owner" for i in range(30)) + "."
    assert len(long_sentence) > M.MAX_TEXT_CHARS
    raw = f"3.7. Response. Response includes the steps taken. {long_sentence}\n3.8. Training. Personnel are trained every year."
    promoted, _ = M.split_candidates({1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}, [])
    assert [p["text"] for p in promoted] == [long_sentence]  # not cut with an ellipsis
    assert raw[promoted[0]["offset"]:].startswith(long_sentence)
    # the old paragraph-level block still shortens a long paragraph for display
    assert M.find_possible_missed({1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": long_sentence}}, [])[0]["source_quote"].endswith(" ...")


def test_a_short_duty_sentence_is_covered_only_by_the_same_extracted_sentence():
    raw = "2.20.19.  Support  MAAs  of  TCAs.  ( T-1 )  Assessment types are described in section 3.3.\n2.20.20. Training. Personnel are trained every year."
    chunks = {1: {"chunk_id": 1, "page_start": 2, "section_title_path": [], "raw_text": raw}}
    assert M.split_candidates(chunks, ["Support MAAs of TCAs."]) == ([], [])  # extracted exactly: not missed
    promoted, rest = M.split_candidates(chunks, ["Support the MAAs of TCAs, as the commander directs."])
    assert [p["text"] for p in promoted] == ["Support MAAs of TCAs."] and rest == []


def test_found_rows_are_placed_by_position_and_an_item_without_one_takes_its_predecessors():
    from services.checklist_service import _merge_in_order
    items = [{"id": "a"}, {"id": "b"}, {"id": "c"}, {"id": "d"}]
    keys = [(1, 10), (1, None), (2, 5), (None, None)]  # b and d have no position of their own
    extra = [{"id": "x"}, {"id": "y"}, {"id": "z"}]
    placed = _merge_in_order(items, keys, extra, [(2, 1), (1, 50), (9, 0)])
    assert [r["id"] for r in placed] == ["a", "b", "y", "x", "c", "d", "z"]  # y (1,50) after b (1,10); x (2,1) before c (2,5); z (9,0) last
    assert _merge_in_order(items, keys, [], []) is items


def test_without_a_chunk_file_there_are_no_candidates_and_no_crash(tmp_path):
    _write(tmp_path, ["The CFP will notify the MCCC."], None)
    result = generate(tmp_path, "doc", "cybersecurity")
    assert result["possible_missed"] == [] and result["summary"]["possible_missed"] == 0


def test_exports_show_the_section_clearly_apart(tmp_path):
    _write(tmp_path, ["The CFP will notify the servicing MCCC within 24 hours of detection."], CHUNKS)
    checklist = generate(tmp_path, "doc", "cybersecurity")
    n_items, n_missed = len(checklist["items"]), len(checklist["possible_missed"])
    assert n_items == 3 and n_missed == 1
    rows = list(csv.DictReader(io.StringIO(to_csv(checklist))))
    assert len(rows) == n_items + n_missed and sum(1 for r in rows if "possible_missed" in r["item_flags"]) == n_missed
    assert sum(1 for r in rows if "found_by_text_scan" in r["item_flags"]) == 2  # the promoted rows are ordinary rows with a flag
    assert "# Possible missed requirements" in to_markdown(checklist)
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(checklist)))["Checklist"]
    assert ws.max_row == 2 + n_items + 1 + 1 + n_missed  # 2 header rows, the items, a blank row, the banner row, the candidates
    banner = [r for r in range(3, ws.max_row + 1) if str(ws.cell(row=r, column=1).value or "").startswith("POSSIBLE MISSED REQUIREMENTS")]
    assert banner == [4 + n_items]
    assert ws.auto_filter.ref.endswith(str(2 + n_items))  # the filter covers the items, found rows included, and not the block
    assert any(ws.cell(row=banner[0] + 1, column=9).coordinate in dv.sqref for dv in ws.data_validations.dataValidation)  # the Status dropdown reaches the candidates too


def test_text_before_the_first_marker_and_bare_bullets_are_units():
    raw = "the Director to notify the CISO within 48 hours.\n2.1. The CFP will log all contacts.\n2.2. Report spills to the NOS."
    units = M.paragraph_units(raw)
    assert units[0].startswith("the Director to notify") and len(units) == 3
    bullets = M.paragraph_units("- Review the access logs every quarter and record the findings.\n- Report anomalies to the security manager.")
    assert len(bullets) == 2
    chunks = {1: {"chunk_id": 1, "page_start": 3, "page_end": 5, "section_title_path": [], "raw_text": "- Review the access logs every quarter and record the findings.\n- Report anomalies to the security manager."}}
    found = M.find_possible_missed(chunks, [])
    assert [f["source_quote"] for f in found] == ["Review the access logs every quarter and record the findings.", "Report anomalies to the security manager."]
    assert found[0]["page_refs"] == [3, 4, 5]  # the whole range the chunk spans

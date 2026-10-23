"""Rough criticality flag (docs/CRITICALITY_FLAG_PLAN.md): sidecar scores, the look-first flag and the exports. The model call is injected, never made."""
import csv
import io
import json

import openpyxl
import pytest

from pipeline.checklist_export import to_csv, to_markdown, to_xlsx
from services import criticality as CR


def _item(i, quote="Commanders shall review logs monthly.", **kw):
    base = {"checklist_item_id": f"CL-{i}", "source_quote": quote, "applies_to": "", "parent_ref": "", "parent_text": "", "passage": f">> {quote} <<", "item_flags": [], "status": "not-started"}
    base.update(kw)
    return base


def _call(table):
    """A fake model: the digit probabilities it returns depend on which row's text is in the prompt."""
    def call(prompt):
        for text, probs in table.items():
            if text in prompt:
                return probs
        return {3: 1.0}
    return call


def test_expected_score_is_the_probability_weighted_digit():
    assert CR.expected_score({4: 0.5, 2: 0.5}) == 3.0
    assert CR.expected_score({5: 1.0}) == 5.0
    assert CR.expected_score({}) is None


def test_expected_score_ignores_other_tokens_and_renormalises():
    assert CR.expected_score({1: 0.1, 5: 0.1, 7: 0.8}) == pytest.approx(3.0)


def test_sidecar_path_sits_beside_the_requirements_file(tmp_path):
    assert CR.sidecar_path(tmp_path / "afi17-203_requirements_gated.jsonl").name == "afi17-203_criticality.jsonl"


def test_rows_are_scored_skipped_rows_get_none_and_the_mean_of_three_wordings_is_stored(tmp_path):
    prompts = []
    items = [_item(1, "Alpha duty."), _item(2, "tail of a sentence", item_flags=["starts_mid_sentence"]), _item(3, "")]
    counts = CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: prompts.append(p) or {4: 0.5, 2: 0.5})
    assert counts["rows"] == 2 and counts["scored"] == 1 and counts["skipped_by_code"] == 1 and len(prompts) == 3
    saved = CR.load(tmp_path / "c.jsonl")
    assert saved["CL-1"]["score"] == 3.0 and len(saved["CL-1"]["scores"]) == 3 and "CL-2" not in saved


def test_unchanged_rows_are_reused_and_a_changed_row_or_model_is_scored_again(tmp_path):
    calls = []
    call = lambda p: calls.append(p) or {3: 1.0}  # noqa: E731
    path = tmp_path / "c.jsonl"
    CR.score_rows([_item(1, "Alpha duty.")], path, call=call)
    counts = CR.score_rows([_item(1, "Alpha duty.")], path, call=call)
    assert counts["reused"] == 1 and len(calls) == 3
    CR.score_rows([_item(1, "Alpha duty, weekly.")], path, call=call)
    assert len(calls) == 6
    CR.score_rows([_item(1, "Alpha duty, weekly.")], path, model="other:model", call=call)
    assert len(calls) == 9


def test_a_service_without_probabilities_stops_the_run_and_writes_no_score(tmp_path):
    def call(prompt):
        raise CR.ProbabilitiesUnavailable("none")
    counts = CR.score_rows([_item(1), _item(2, "Second duty.")], tmp_path / "c.jsonl", call=call)
    assert counts["unavailable"] and counts["scored"] == 0 and CR.load(tmp_path / "c.jsonl") == {}


def test_an_unreachable_server_stops_after_five_failures_in_a_row_and_keeps_earlier_rows(tmp_path):
    def call(prompt):
        if "Good duty." not in prompt:
            raise OSError("down")
        return {5: 1.0}
    items = [_item(0, "Good duty.")] + [_item(i, f"Bad duty {i}.") for i in range(1, 9)]
    counts = CR.score_rows(items, tmp_path / "c.jsonl", call=call)
    assert counts["aborted"] and counts["scored"] == 1 and counts["errors"] == 5
    assert set(CR.load(tmp_path / "c.jsonl")) == {"CL-0"}


def test_rows_get_a_band_by_rank_and_no_number(tmp_path):
    items = [_item(i, f"Duty number {i}.") for i in range(20)]
    table = {f"Duty number {i}.": {1: 1.0 - i / 20, 5: i / 20} for i in range(20)}
    CR.score_rows(items, tmp_path / "c.jsonl", call=_call(table))
    assert CR.apply(items, tmp_path / "c.jsonl") == 20
    bands = [i["criticality"] for i in sorted(items, key=lambda i: -int(i["checklist_item_id"][3:]))]
    assert bands == ["Critical"] + ["High"] * 4 + ["Medium"] * 7 + ["Low"] * 8
    assert not any("criticality_score" in i or "criticality_flag" in i for i in items)


def test_a_single_scored_row_is_critical_not_empty(tmp_path):
    items = [_item(1, "Alpha duty.")]
    CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: {3: 1.0})
    CR.apply(items, tmp_path / "c.jsonl")
    assert items[0]["criticality"] == "Critical"


def test_with_no_sidecar_or_a_stale_record_nothing_is_added(tmp_path):
    items = [_item(1, "Alpha duty.")]
    assert CR.apply(items, tmp_path / "missing.jsonl") == 0 and "criticality" not in items[0]
    CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: {5: 1.0})
    changed = [_item(1, "Alpha duty changed.")]
    assert CR.apply(changed, tmp_path / "c.jsonl") == 0 and "criticality" not in changed[0]


def test_ollama_reply_without_probabilities_raises_unavailable(monkeypatch):
    class R:
        def raise_for_status(self): pass
        def json(self): return {"message": {"content": "3"}}
    monkeypatch.setattr("requests.post", lambda *a, **k: R())
    with pytest.raises(CR.ProbabilitiesUnavailable):
        CR.ollama_digit_probs("p", "m", "http://x")


def test_ollama_reply_with_probabilities_gives_digit_probabilities(monkeypatch):
    import math

    class R:
        def raise_for_status(self): pass
        def json(self): return {"logprobs": [{"top_logprobs": [{"token": "4", "logprob": math.log(0.6)}, {"token": " 3", "logprob": math.log(0.3)}, {"token": "The", "logprob": math.log(0.1)}]}]}
    monkeypatch.setattr("requests.post", lambda *a, **k: R())
    probs = CR.ollama_digit_probs("p", "m", "http://x")
    assert round(probs[4], 3) == 0.6 and round(probs[3], 3) == 0.3 and set(probs) == {3, 4}


# --- exports: unchanged without scores, one extra column with them ---

def _env(items):
    return {"format": "reqbot-checklist", "format_version": "1.1", "generated_at": "2026-10-17T00:00:00+00:00", "generator": {"tool": "reqbot"}, "profile": "cybersecurity",
            "summary": {"total_items": len(items), "items_requiring_review": 0}, "items": items}


def test_exports_without_scores_have_no_criticality_column():
    env = _env([_item(1)])
    assert "criticality" not in to_csv(env)
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(env))).active
    assert "Criticality (rough)" not in [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    assert "Rough criticality" not in to_markdown(env)


def test_exports_with_bands_show_the_word_and_no_number():
    env = _env([_item(1, criticality="Critical"), _item(2, "Second duty.", criticality="Low")])
    rows = list(csv.DictReader(io.StringIO(to_csv(env))))
    assert (rows[0]["criticality"], rows[1]["criticality"]) == ("Critical", "Low")
    ws = openpyxl.load_workbook(io.BytesIO(to_xlsx(env))).active
    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    col = headers.index("Criticality (rough)") + 1
    assert headers[col - 2] == "Audit Question (draft)" and ws.cell(row=5, column=col).value == "Critical" and ws.cell(row=6, column=col).value == "Low"
    assert "**Rough criticality:** Critical" in to_markdown(env)


def test_the_document_name_goes_into_the_prompt_and_a_record_with_another_context_is_stale(tmp_path):
    prompts = []
    items = [_item(1, "Alpha duty.")]
    CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: prompts.append(p) or {3: 1.0}, document="AFI 17-203")
    assert all("requirement from the publication AFI 17-203" in p for p in prompts)
    assert CR.apply(items, tmp_path / "c.jsonl") == 1  # the record keeps its own context, so it still matches
    counts = CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: {3: 1.0}, document="AFI 99-999")
    assert counts["reused"] == 0 and counts["scored"] == 1


def test_a_row_where_one_wording_fails_is_not_stored_as_a_partial_mean_and_is_retried(tmp_path):
    calls = []

    def flaky(prompt):
        calls.append(prompt)
        if len(calls) == 3:
            raise OSError("timed out")
        return {4: 1.0}
    path = tmp_path / "c.jsonl"
    counts = CR.score_rows([_item(1, "Alpha duty.")], path, call=flaky)
    assert counts["scored"] == 0 and counts["errors"] == 1 and CR.load(path) == {}
    counts = CR.score_rows([_item(1, "Alpha duty.")], path, call=lambda p: {4: 1.0})
    assert counts["scored"] == 1 and CR.load(path)["CL-1"]["scores"] == [4.0, 4.0, 4.0]


def test_ollama_reply_with_malformed_token_entries_raises_unavailable(monkeypatch):
    class R:
        def raise_for_status(self): pass
        def json(self): return {"logprobs": [{"top_logprobs": [{"token": "4"}, "junk"]}]}
    monkeypatch.setattr("requests.post", lambda *a, **k: R())
    with pytest.raises(CR.ProbabilitiesUnavailable):
        CR.ollama_digit_probs("p", "m", "http://x")


def test_rows_the_questions_job_withheld_get_no_score_and_no_flag(tmp_path):
    from services import audit_questions as AQ
    items = [_item(1, "Alpha duty."), _item(2, "Beta permission."), _item(3, "Gamma duty.")]
    AQ.draft_questions(items, tmp_path / "q.jsonl", call=lambda p: json.dumps({"question": None if "Beta" in p else "Do they do it?"}))
    withheld = CR.withheld_ids(tmp_path / "q.jsonl")
    assert withheld == {"CL-2"}
    counts = CR.score_rows(items, tmp_path / "c.jsonl", call=lambda p: {4: 1.0}, skip_ids=withheld)
    assert counts["scored"] == 2 and set(CR.load(tmp_path / "c.jsonl")) == {"CL-1", "CL-3"}
    CR.score_rows(items, tmp_path / "c2.jsonl", call=lambda p: {4: 1.0})  # scored before the questions result was known
    CR.apply(items, tmp_path / "c2.jsonl", skip_ids=withheld)
    assert "criticality" not in items[1] and "criticality" in items[0]

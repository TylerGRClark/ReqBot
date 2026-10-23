"""Phase 23 integration gate (WP-23.6).

Verifies Phase 23 features work end-to-end across module boundaries.
No live Qdrant/Ollama required; no frontend assertions (WP-23.1 scroll is
manual browser verification).

Features exercised:
  WP-23.2 — XLSX export via checklist service → to_xlsx()
  WP-23.4 — cybersecurity profile has skip_sections; param accepted by the
             structure-aware chunk_text entry point
  Regression — CSV/JSON/MD checklist generation unchanged by Phase 23

WP-23.3's overlap-guard/page-contiguity tests were removed under WP-34.1 (legacy
pymupdf/pdfplumber chunking, the only thing those functions applied to, was
deleted -- see archive/PHASE34_REQUIREMENTS.md). WP-23.5's recovered_truncated tests were removed
with the list-writing extraction step it belonged to (replaced by clause labeling, which writes no
truncated lists to recover).
"""
import io
import json
import shutil
from pathlib import Path

import openpyxl
import pytest

from core.profiles import load_profile
from pipeline.checklist_export import to_csv, to_json, to_markdown, to_xlsx
from services.checklist_service import generate

_FIXTURE = Path(__file__).resolve().parent.parent / "fixtures" / "sample_normalized_reqs.jsonl"
DOC_KEY = "SYNTHETIC_TEST_DOC"


@pytest.fixture()
def processed_dir(tmp_path: Path) -> Path:
    run_dir = tmp_path / f"{DOC_KEY}_20260101_000000"
    run_dir.mkdir()
    shutil.copy(_FIXTURE, run_dir / f"{DOC_KEY}_requirements_normalized.jsonl")
    return tmp_path


# ---------------------------------------------------------------------------
# WP-23.2 — XLSX export integration (service → workbook)
# ---------------------------------------------------------------------------

def test_xlsx_from_service_produces_valid_workbook(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    raw = to_xlsx(checklist)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    assert wb is not None


def test_xlsx_item_count_matches_checklist(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    raw = to_xlsx(checklist)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    data_rows = ws.max_row - 4  # the title, facts and two header rows
    assert data_rows == len(checklist["items"])


def test_xlsx_has_frozen_pane(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    raw = to_xlsx(checklist)
    wb = openpyxl.load_workbook(io.BytesIO(raw))
    ws = wb.active
    assert ws.freeze_panes is not None


def test_xlsx_api_route_format_accepted(processed_dir):
    # Smoke: the 'xlsx' format string is wired into the route layer.
    # Import the route module to confirm it doesn't error on import.
    from api.routes.checklist import router  # noqa: F401


# ---------------------------------------------------------------------------
# WP-23.4 — skip_sections in cybersecurity profile + param threading
# ---------------------------------------------------------------------------

def test_cybersecurity_profile_has_skip_sections():
    profile = load_profile("cybersecurity")
    assert "skip_sections" in profile
    skip = profile["skip_sections"]
    assert isinstance(skip, list)
    assert len(skip) > 0


def test_cybersecurity_profile_skip_sections_contains_expected_entries():
    profile = load_profile("cybersecurity")
    skip_upper = [s.upper() for s in profile["skip_sections"]]
    for expected in ("GLOSSARY", "REFERENCES", "ACRONYMS"):
        assert expected in skip_upper, f"Expected '{expected}' in skip_sections"


def test_chunk_text_run_structure_aware_accepts_skip_sections_param():
    import inspect

    import pipeline.chunk_text as ct
    sig = inspect.signature(ct.run_structure_aware)
    assert "skip_sections" in sig.parameters


# ---------------------------------------------------------------------------
# Regression — Phase 21/22 checklist formats unchanged
# ---------------------------------------------------------------------------

def test_csv_generation_regression(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    result = to_csv(checklist)
    assert isinstance(result, str)
    assert "source_quote" in result


def test_json_generation_regression(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    result = to_json(checklist)
    data = json.loads(result)
    assert "items" in data
    assert len(data["items"]) > 0


def test_markdown_generation_regression(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    result = to_markdown(checklist)
    assert isinstance(result, str)
    assert "#" in result


def test_all_four_export_formats_produce_output(processed_dir):
    checklist = generate(processed_dir, DOC_KEY, "cybersecurity")
    assert to_csv(checklist)
    assert to_json(checklist)
    assert to_markdown(checklist)
    assert to_xlsx(checklist)
